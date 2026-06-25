from pathlib import Path
import argparse
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from card_tracker.report_paths import timestamped_report_dir


HEADERS = [
    "Quantity",
    "Card Name",
    "Card Number",
    "Holo Type",
    "Language",
    "Condition",
    "Grade Company",
    "Grade Received",
    "Notes",
    "Matched Card ID",
    "Resolution Status",
    "Matched Set",
    "Matched Set Code",
    "TCGcollector Card ID",
    "Card Detail URL",
]


EXAMPLE_ROWS = [
    [1, "Drapion V", "118/196", "", "English", "NM/Mint", "", "", "example row; replace me"],
    [1, "", "", "Reverse Holo", "English", "LP (Light)", "", "", ""],
    [1, "", "", "Energy Holo", "Japanese", "NM/Mint", "", "", ""],
    [1, "", "", "Poke Ball Holo", "Korean", "NM/Mint", "", "", "currently unsupported by the TCGcollector importer"],
]


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Create an editable binder intake workbook.")
    # Argument examples:
    #   python scripts/build_binder_intake_workbook.py
    #   python scripts/build_binder_intake_workbook.py --output reports/binder_intake/YYYYMMDD_HHMMSS_label/my_binder.xlsx
    parser.add_argument("--output", type=Path, help="Workbook path to create.")
    parser.add_argument(
        "--report-label",
        default="binder_intake_template",
        help="Label for the generated timestamped report folder when --output is omitted.",
    )
    parser.add_argument("--rows", type=int, default=250, help="Number of editable intake rows.")
    return parser.parse_args()


def resolve_output(path: Path) -> Path:
    return path if path.is_absolute() else ROOT / path


def main() -> None:
    args = parse_args()
    output = resolve_output(args.output) if args.output else timestamped_report_dir("binder_intake", args.report_label) / "binder_intake_template.xlsx"
    output.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Binder Intake"
    sheet.freeze_panes = "A2"

    for column, header in enumerate(HEADERS, start=1):
        cell = sheet.cell(1, column, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F2937")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_number in range(2, args.rows + 2):
        row_values = EXAMPLE_ROWS[row_number - 2] if row_number - 2 < len(EXAMPLE_ROWS) else [None] * 6
        for column, value in enumerate(row_values, start=1):
            sheet.cell(row_number, column, value)

    language_validation = DataValidation(
        type="list",
        formula1='"English,Japanese,Simplified Chinese,Korean"',
        allow_blank=False,
    )
    condition_validation = DataValidation(
        type="list",
        formula1='"NM/Mint,LP (Light),MP (Moderate),HP (Heavy),DMG (Damaged),Graded"',
        allow_blank=False,
    )
    holo_validation = DataValidation(
        type="list",
        formula1='"Reverse Holo,Energy Holo,Poke Ball Holo"',
        allow_blank=True,
    )
    grading_company_validation = DataValidation(
        type="list",
        formula1='"PSA,BGS,CGC,SGC,TAG,ACE,PCG"',
        allow_blank=True,
    )
    grade_received_validation = DataValidation(
        type="list",
        formula1='"1,1.5,2,2.5,3,3.5,4,4.5,5,5.5,6,6.5,7,7.5,8,8.5,9,9.5,10"',
        allow_blank=True,
    )
    sheet.add_data_validation(language_validation)
    sheet.add_data_validation(condition_validation)
    sheet.add_data_validation(holo_validation)
    sheet.add_data_validation(grading_company_validation)
    sheet.add_data_validation(grade_received_validation)
    holo_validation.add(f"D2:D{args.rows + 1}")
    language_validation.add(f"E2:E{args.rows + 1}")
    condition_validation.add(f"F2:F{args.rows + 1}")
    grading_company_validation.add(f"G2:G{args.rows + 1}")
    grade_received_validation.add(f"H2:H{args.rows + 1}")

    widths = {
        "A": 10,
        "B": 26,
        "C": 16,
        "D": 18,
        "E": 20,
        "F": 16,
        "G": 16,
        "H": 16,
        "I": 36,
        "J": 16,
        "K": 24,
        "L": 28,
        "M": 18,
        "N": 20,
        "O": 52,
    }
    for column_letter, width in widths.items():
        sheet.column_dimensions[column_letter].width = width

    for row in sheet.iter_rows(min_row=2, max_row=args.rows + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{args.rows + 1}"

    meta = workbook.create_sheet("Instructions")
    meta["A1"] = "Binder intake workflow"
    meta["A1"].font = Font(bold=True, size=14)
    meta["A3"] = "Fill Quantity, Card Name, Card Number, optional Holo Type, Language, Condition, and optional Notes."
    meta["A4"] = "Leave the matched/result columns blank; scripts/import_binder_inventory.py fills them during resolution."
    meta["A5"] = "Card Number should be the printed collector number, such as 118/196 or GG49/GG70."
    meta["A6"] = "Use Grade Company and Grade Received only when Condition is Graded."
    meta["A7"] = "Korean is included for tracking, but current automated set import is limited to English, Japanese, and Simplified Chinese."
    meta["A8"] = "Leave Holo Type blank for the default card pattern. Select Reverse Holo, Energy Holo, or Poke Ball Holo only for parallel variants."
    meta.column_dimensions["A"].width = 120

    workbook.save(output)
    print(f"Created binder intake workbook: {output}")


if __name__ == "__main__":
    main()
