from pathlib import Path
import argparse
import sqlite3
from typing import Any

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "data" / "card_tracker.sqlite"
DEFAULT_NOTES = "inventory_workbook_import"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Import editable inventory counts from a protected inventory workbook."
    )
    # Argument examples:
    #   reports/inventory_workbook/YYYYMMDD_HHMMSS_cbb5c/cbb5c_gem_pack_vol_5_inventory.xlsx
    #   reports/inventory_workbook/YYYYMMDD_HHMMSS_cbb5c/cbb5c_gem_pack_vol_5_inventory.xlsx --dry-run
    #   reports/inventory_workbook/YYYYMMDD_HHMMSS_cbb5c/cbb5c_gem_pack_vol_5_inventory.xlsx --condition raw
    parser.add_argument("workbook", help="Path to the inventory workbook to import.")
    parser.add_argument("--sheet", default="Inventory", help="Workbook sheet name to read.")
    parser.add_argument("--header-row", type=int, default=4, help="Header row number.")
    parser.add_argument("--condition", default="raw", help="Condition value for imported rows.")
    parser.add_argument("--notes", default=DEFAULT_NOTES, help="Notes marker for imported rows.")
    parser.add_argument(
        "--import-label",
        default="",
        help="Inventory import batch label. Defaults to the workbook filename stem.",
    )
    parser.add_argument("--dry-run", action="store_true", help="Validate and summarize without writing.")
    return parser.parse_args()


def workbook_path(value: str) -> Path:
    path = Path(value)
    if not path.is_absolute():
        path = ROOT / path
    return path


def count_from_cell(value: Any, row_number: int) -> int:
    if value in (None, ""):
        return 0

    if isinstance(value, float) and value.is_integer():
        value = int(value)

    if not isinstance(value, int):
        raise ValueError(f"Row {row_number}: Count must be a whole number, got {value!r}.")

    if value < 0:
        raise ValueError(f"Row {row_number}: Count must be 0 or greater, got {value}.")

    return value


def read_inventory_rows(path: Path, sheet_name: str, header_row: int) -> dict[int, int]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet {sheet_name!r} was not found in {path}.")

    sheet = workbook[sheet_name]
    headers = {
        sheet.cell(header_row, column).value: column
        for column in range(1, sheet.max_column + 1)
    }

    missing_headers = {"Card ID", "Count"} - set(headers)
    if missing_headers:
        raise ValueError(f"Missing required header(s): {', '.join(sorted(missing_headers))}.")

    card_id_column = headers["Card ID"]
    count_column = headers["Count"]
    counts_by_card_id: dict[int, int] = {}

    for row_number in range(header_row + 1, sheet.max_row + 1):
        card_id = sheet.cell(row_number, card_id_column).value
        if card_id in (None, ""):
            continue

        if isinstance(card_id, float) and card_id.is_integer():
            card_id = int(card_id)

        if not isinstance(card_id, int):
            raise ValueError(f"Row {row_number}: Card ID must be a whole number, got {card_id!r}.")

        if card_id in counts_by_card_id:
            raise ValueError(f"Duplicate Card ID in workbook: {card_id}.")

        count = count_from_cell(sheet.cell(row_number, count_column).value, row_number)
        counts_by_card_id[card_id] = count

    if not counts_by_card_id:
        raise ValueError("No card rows were found in the workbook.")

    return counts_by_card_id


def validate_card_ids(conn: sqlite3.Connection, card_ids: list[int]) -> None:
    placeholders = ",".join("?" for _ in card_ids)
    found_ids = {
        row[0]
        for row in conn.execute(
            f"SELECT id FROM cards WHERE id IN ({placeholders})",
            card_ids,
        )
    }
    missing_ids = sorted(set(card_ids) - found_ids)
    if missing_ids:
        preview = ", ".join(str(card_id) for card_id in missing_ids[:20])
        raise ValueError(f"Workbook contains card IDs that are not in the database: {preview}.")


def inventory_import_id(
    conn: sqlite3.Connection,
    import_label: str,
    source_path: Path,
    notes: str,
) -> int:
    conn.execute(
        """
        CREATE TABLE IF NOT EXISTS inventory_imports (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            import_label TEXT NOT NULL UNIQUE,
            import_type TEXT NOT NULL,
            source_name TEXT,
            source_path TEXT,
            notes TEXT,
            imported_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
            created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
        )
        """
    )
    conn.execute(
        """
        INSERT INTO inventory_imports (import_label, import_type, source_name, source_path, notes)
        VALUES (?, 'inventory_workbook', 'inventory_workbook', ?, ?)
        ON CONFLICT(import_label) DO UPDATE SET
            source_path = excluded.source_path,
            notes = excluded.notes
        """,
        (import_label, str(source_path), notes),
    )
    return int(conn.execute("SELECT id FROM inventory_imports WHERE import_label = ?", (import_label,)).fetchone()[0])


def import_inventory(
    counts_by_card_id: dict[int, int],
    condition: str,
    notes: str,
    import_label: str,
    source_path: Path,
    dry_run: bool,
) -> tuple[int, int, int, int]:
    card_ids = sorted(counts_by_card_id)
    nonzero_items = [(card_id, count) for card_id, count in counts_by_card_id.items() if count > 0]

    with sqlite3.connect(DB_PATH) as conn:
        validate_card_ids(conn, card_ids)

        placeholders = ",".join("?" for _ in card_ids)
        rows_to_replace = conn.execute(
            f"""
            SELECT COUNT(*)
            FROM card_inventory
            WHERE card_id IN ({placeholders})
                AND (sale_status = 'reference' OR notes = ?)
            """,
            [*card_ids, notes],
        ).fetchone()[0]

        if dry_run:
            return len(card_ids), len(nonzero_items), sum(count for _, count in nonzero_items), rows_to_replace

        with conn:
            batch_id = inventory_import_id(conn, import_label, source_path, notes)
            conn.execute(
                f"""
                DELETE FROM card_inventory
                WHERE card_id IN ({placeholders})
                    AND (sale_status = 'reference' OR notes = ?)
                """,
                [*card_ids, notes],
            )

            conn.executemany(
                """
                INSERT INTO card_inventory (
                    card_id,
                    import_id,
                    condition,
                    quantity,
                    cost_basis_cents,
                    acquisition_date,
                    sale_status,
                    notes
                )
                VALUES (?, ?, ?, ?, 0, NULL, 'inventory', ?)
                """,
                [(card_id, batch_id, condition, count, notes) for card_id, count in nonzero_items],
            )

    return len(card_ids), len(nonzero_items), sum(count for _, count in nonzero_items), rows_to_replace


def main() -> None:
    args = parse_args()
    path = workbook_path(args.workbook)
    import_label = args.import_label.strip() or path.stem
    counts_by_card_id = read_inventory_rows(path, args.sheet, args.header_row)
    card_rows, imported_rows, imported_quantity, replaced_rows = import_inventory(
        counts_by_card_id=counts_by_card_id,
        condition=args.condition,
        notes=args.notes,
        import_label=import_label,
        source_path=path,
        dry_run=args.dry_run,
    )

    action = "Would import" if args.dry_run else "Imported"
    print(f"{action} {imported_rows} inventory rows totaling {imported_quantity} cards.")
    print(f"Validated {card_rows} workbook card rows.")
    print(f"Replaced {replaced_rows} existing reference/workbook-import rows.")


if __name__ == "__main__":
    main()
