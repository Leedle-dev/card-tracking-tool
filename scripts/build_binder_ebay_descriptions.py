from __future__ import annotations

from pathlib import Path
import argparse
import sqlite3
import sys

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from card_tracker.reporting import ReportBuilder, SheetSpec, WorkbookSpec


DB_PATH = ROOT / "data" / "card_tracker.sqlite"

DESCRIPTION_TEMPLATE = """You are purchasing the exact {POKEMON} card shown in the photos.

Card:
{POKEMON_NAME} - {CARD_NUMBER}
Set: {SET_NAME} / {SET_CODE}
Language: {LANGUAGE}
Condition: {CONDITION}

The card pictured is the exact card you will receive. Please review the photos carefully for condition, surface, edge, corner, and centering details.

I photograph cards individually so buyers can see the actual item before purchase. I try to display any visible wear, whitening, scratches, or print lines.

Shipping:
Card will be sleeved, placed in a top loader or semi-rigid holder, and shipped securely.

Combined shipping is available! Check out my shop for additional cards to add to your cart before checkout and get a discount!

Questions are welcome. Thanks for checking out Lee's TCG & Collectables!"""


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build ready-to-copy eBay descriptions for binder inventory cards.")
    # Argument examples:
    #   python scripts/build_binder_ebay_descriptions.py
    #   python scripts/build_binder_ebay_descriptions.py --output-dir reports/binder_ebay_descriptions/test
    parser.add_argument("--db-path", type=Path, default=DB_PATH, help="Path to the local SQLite database.")
    parser.add_argument("--output-dir", type=Path, help="Optional explicit report output directory.")
    parser.add_argument("--report-label", default="binder_inventory", help="Label for the timestamped report folder.")
    parser.add_argument(
        "--notes-prefix",
        default="binder_inventory_import",
        help="card_inventory.notes prefix used to identify binder inventory rows.",
    )
    parser.add_argument(
        "--import-label",
        default="",
        help="Optional inventory_imports.import_label filter, such as binder_intake_run_two.",
    )
    parser.add_argument(
        "--order-workbook",
        type=Path,
        help="Optional binder intake workbook whose row order should drive report row order.",
    )
    parser.add_argument(
        "--order-sheet",
        default="Binder Intake",
        help="Sheet name to read when --order-workbook is supplied.",
    )
    return parser.parse_args()


def resolve_path(path: Path) -> Path:
    return path if path.is_absolute() else ROOT / path


def fetch_binder_inventory(db_path: Path, notes_prefix: str, import_label: str = "") -> list[sqlite3.Row]:
    conn = sqlite3.connect(db_path)
    conn.row_factory = sqlite3.Row
    try:
        import_filter = "AND ii.import_label = ?" if import_label else "AND ci.notes LIKE ?"
        param = import_label if import_label else f"{notes_prefix}%"
        return conn.execute(
            f"""
            SELECT
                ci.id AS inventory_id,
                ci.quantity,
                ci.condition,
                ci.notes,
                ii.import_label,
                c.id AS card_id,
                c.name AS card_name,
                COALESCE(c.pokemon_name, c.name) AS pokemon_name,
                c.card_number,
                c.holo_pattern,
                c.rarity,
                sc.set_name,
                sc.set_code,
                sc.language
            FROM card_inventory ci
            LEFT JOIN inventory_imports ii ON ii.id = ci.import_id
            JOIN cards c ON c.id = ci.card_id
            JOIN set_catalog sc ON sc.id = c.set_catalog_id
            WHERE ci.sale_status = 'inventory'
                AND ci.quantity > 0
                {import_filter}
            ORDER BY sc.language, sc.set_name, c.card_number, c.name, ci.id
            """,
            (param,),
        ).fetchall()
    finally:
        conn.close()


def blank_if_none(value: object) -> str:
    return "" if value is None else str(value)


def description_for(row: sqlite3.Row) -> str:
    card_name = blank_if_none(row["card_name"])
    holo_pattern = blank_if_none(row["holo_pattern"])
    card_line_name = f"{card_name} ({holo_pattern})" if holo_pattern else card_name
    return DESCRIPTION_TEMPLATE.format(
        POKEMON=card_name,
        POKEMON_NAME=card_line_name,
        CARD_NUMBER=blank_if_none(row["card_number"]),
        SET_NAME=blank_if_none(row["set_name"]),
        SET_CODE=blank_if_none(row["set_code"]),
        LANGUAGE=blank_if_none(row["language"]),
        CONDITION=blank_if_none(row["condition"]),
    )


def ebay_title_for(row: sqlite3.Row) -> str:
    parts = [
        blank_if_none(row["card_name"]),
        blank_if_none(row["card_number"]),
        blank_if_none(row["holo_pattern"]),
        blank_if_none(row["set_name"]),
        blank_if_none(row["set_code"]),
        blank_if_none(row["language"]),
        "Pokemon Card",
    ]
    title = " ".join(part for part in parts if part)
    condition = blank_if_none(row["condition"])
    return f"{title} - {condition}" if condition else title


def workbook_order_map(path: Path, sheet_name: str) -> dict[int, int]:
    workbook_path = resolve_path(path)
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise SystemExit(f"Sheet {sheet_name!r} was not found in {workbook_path}.")
    sheet = workbook[sheet_name]
    headers = {
        str(sheet.cell(1, column).value or "").strip(): column
        for column in range(1, sheet.max_column + 1)
    }
    card_id_column = headers.get("Matched Card ID")
    if not card_id_column:
        raise SystemExit(f"{workbook_path} does not have a 'Matched Card ID' column.")
    order = {}
    for row_number in range(2, sheet.max_row + 1):
        value = sheet.cell(row_number, card_id_column).value
        if value in (None, ""):
            continue
        if isinstance(value, float) and value.is_integer():
            value = int(value)
        try:
            card_id = int(value)
        except (TypeError, ValueError):
            continue
        order.setdefault(card_id, row_number)
    return order


def build_rows(source_rows: list[sqlite3.Row], order: dict[int, int] | None = None) -> list[dict[str, object]]:
    output_rows = []
    for row in source_rows:
        output_rows.append(
            {
                "binder_row": order.get(int(row["card_id"]), "") if order else "",
                "inventory_id": row["inventory_id"],
                "card_id": row["card_id"],
                "quantity": row["quantity"],
                "ebay_title": ebay_title_for(row),
                "card_name": blank_if_none(row["card_name"]),
                "card_number": blank_if_none(row["card_number"]),
                "holo_pattern": blank_if_none(row["holo_pattern"]),
                "set_name": blank_if_none(row["set_name"]),
                "set_code": blank_if_none(row["set_code"]),
                "language": blank_if_none(row["language"]),
                "condition": blank_if_none(row["condition"]),
                "rarity": blank_if_none(row["rarity"]),
                "import_label": blank_if_none(row["import_label"]),
                "inventory_notes": blank_if_none(row["notes"]),
                "ebay_description": description_for(row),
            }
        )
    return sorted(output_rows, key=lambda row: (row["binder_row"] or 999999, row["card_id"]))


def main() -> None:
    args = parse_args()
    db_path = resolve_path(args.db_path)
    source_rows = fetch_binder_inventory(db_path, args.notes_prefix, args.import_label)
    if not source_rows:
        raise SystemExit("No binder inventory rows found.")

    order = workbook_order_map(args.order_workbook, args.order_sheet) if args.order_workbook else None
    rows = build_rows(source_rows, order)
    fieldnames = [
        "binder_row",
        "inventory_id",
        "card_id",
        "quantity",
        "ebay_title",
        "card_name",
        "card_number",
        "holo_pattern",
        "set_name",
        "set_code",
        "language",
        "condition",
        "rarity",
        "import_label",
        "inventory_notes",
        "ebay_description",
    ]

    report = ReportBuilder(
        report_type="binder_ebay_descriptions",
        label=args.report_label,
        output_dir=args.output_dir,
        metadata={"notes_prefix": args.notes_prefix, "import_label": args.import_label, "row_count": len(rows)},
    )
    report.write_tsv("binder_ebay_descriptions.tsv", rows, fieldnames=fieldnames)
    report.write_workbook(
        WorkbookSpec(
            filename="binder_ebay_descriptions.xlsx",
            sheets=[
                SheetSpec(
                    name="Descriptions",
                    rows=rows,
                    fieldnames=fieldnames,
                    freeze_panes="A2",
                    column_widths={
                        "inventory_id": 12,
                        "binder_row": 10,
                        "card_id": 10,
                        "quantity": 10,
                        "ebay_title": 56,
                        "card_name": 26,
                        "card_number": 16,
                        "holo_pattern": 18,
                        "set_name": 26,
                        "set_code": 14,
                        "language": 16,
                        "condition": 16,
                        "rarity": 20,
                        "import_label": 24,
                        "inventory_notes": 32,
                        "ebay_description": 84,
                    },
                    integer_columns=["inventory_id", "card_id", "quantity"],
                )
            ],
        )
    )
    report.write_metadata()
    print(f"Wrote {len(rows)} eBay descriptions.")
    print(report.output_dir)


if __name__ == "__main__":
    main()
