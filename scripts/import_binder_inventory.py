from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from urllib.parse import quote_plus
import argparse
import re
import sqlite3
import sys
import unicodedata

from openpyxl import load_workbook
from sqlalchemy import select

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from card_tracker.db.models import Card, CardInventory, GradingCompany, InventoryImport, SetCatalog
from card_tracker.db.session import create_session_factory
from import_tcgcollector_set import (
    BASE_URL,
    LANGUAGE_BY_REGION,
    clear_existing_local_images,
    ensure_schema_columns,
    extract_cards,
    fetch_url_bytes,
    image_dir_for_set,
    import_card,
    normalize_region,
    resolve_set,
)


DB_PATH = ROOT / "data" / "card_tracker.sqlite"
DEFAULT_NOTES = "binder_inventory_import"
SUPPORTED_TCGCOLLECTOR_LANGUAGES = {"English", "Japanese", "Simplified Chinese"}
CONDITION_ALIASES = {
    "nm mint": "NM/Mint",
    "near mint": "NM/Mint",
    "mint": "NM/Mint",
    "raw": "NM/Mint",
    "lp light": "LP (Light)",
    "light": "LP (Light)",
    "mp moderate": "MP (Moderate)",
    "moderate": "MP (Moderate)",
    "hp heavy": "HP (Heavy)",
    "heavy": "HP (Heavy)",
    "dmg damaged": "DMG (Damaged)",
    "damaged": "DMG (Damaged)",
    "graded": "Graded",
}
HOLO_TYPE_ALIASES = {
    "": "",
    "default": "",
    "normal": "",
    "standard": "",
    "regular": "",
    "non holo": "",
    "nonholo": "",
    "reverse": "Reverse Holo",
    "reverse holo": "Reverse Holo",
    "reverse foil": "Reverse Holo",
    "energy": "Energy Holo",
    "energy holo": "Energy Holo",
    "energy foil": "Energy Holo",
    "poke ball": "Poke Ball Holo",
    "pokeball": "Poke Ball Holo",
    "poke ball holo": "Poke Ball Holo",
    "pokeball holo": "Poke Ball Holo",
    "poke ball foil": "Poke Ball Holo",
}

INPUT_HEADERS = ["Quantity", "Card Name", "Card Number", "Language", "Condition", "Notes"]
OPTIONAL_INPUT_HEADERS = ["Grade Company", "Grade Received", "Holo Type"]
HEADER_ALIASES = {
    "NM/Mint": "Condition",
    "LP (Light)": "Condition",
    "MP (Moderate)": "Condition",
    "HP (Heavy)": "Condition",
    "DMG (Damaged)": "Condition",
    "Graded": "Condition",
}
RESULT_HEADERS = [
    "Matched Card ID",
    "Resolution Status",
    "Matched Set",
    "Matched Set Code",
    "Matched Holo Type",
    "TCGcollector Card ID",
    "Card Detail URL",
]


@dataclass
class IntakeRow:
    row_number: int
    quantity: int
    card_name: str
    card_number: str
    holo_pattern: str | None
    language: str
    condition: str
    grade_company: str | None
    grade_received: float | None
    notes: str | None


@dataclass
class ResolvedCard:
    card_id: int
    name: str
    card_number: str | None
    holo_pattern: str | None
    set_name: str | None
    set_code: str | None
    language: str | None
    tcgcollector_card_id: int | None
    card_detail_url: str | None


@dataclass
class RemoteCardMatch:
    name: str
    card_number: str
    set_name: str
    tcgcollector_card_id: int
    card_detail_url: str


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Resolve a binder intake workbook, optionally import missing TCGcollector sets, and update inventory."
    )
    # Argument examples:
    #   python scripts/import_binder_inventory.py reports/binder_intake/YYYYMMDD_HHMMSS_label/binder_intake_template.xlsx
    #   python scripts/import_binder_inventory.py reports/binder_intake/YYYYMMDD_HHMMSS_label/binder_intake_template.xlsx --write-results
    #   python scripts/import_binder_inventory.py reports/binder_intake/YYYYMMDD_HHMMSS_label/binder_intake_template.xlsx --import-missing-sets --apply --write-results
    parser.add_argument("workbook", help="Path to the binder intake workbook.")
    parser.add_argument("--sheet", default="Binder Intake", help="Workbook sheet name to read.")
    parser.add_argument("--db-path", type=Path, default=DB_PATH, help="SQLite database path.")
    parser.add_argument("--apply", action="store_true", help="Write inventory changes. Default is dry run.")
    parser.add_argument("--write-results", action="store_true", help="Write resolution result columns back to the workbook.")
    parser.add_argument(
        "--import-missing-sets",
        action="store_true",
        help="Import supported missing sets found via TCGcollector card search.",
    )
    parser.add_argument(
        "--replace-reference",
        action="store_true",
        help="Allow binder quantities to replace existing reference inventory rows.",
    )
    parser.add_argument(
        "--merge-existing",
        action="store_true",
        help="Add binder quantity to an existing non-reference inventory row instead of reporting a conflict.",
    )
    parser.add_argument("--notes", default=DEFAULT_NOTES, help="Notes marker for inventory rows created by this import.")
    parser.add_argument(
        "--import-label",
        default="",
        help="Inventory import batch label. Defaults to the workbook filename stem.",
    )
    return parser.parse_args()


def resolve_path(path: Path | str) -> Path:
    value = Path(path)
    return value if value.is_absolute() else ROOT / value


def normalize_text(value: str | None) -> str:
    if not value:
        return ""
    value = normalize_search_query(value)
    value = unicodedata.normalize("NFKD", value)
    value = value.encode("ascii", "ignore").decode("ascii")
    value = value.lower()
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return " ".join(value.split())


def normalize_search_query(value: str) -> str:
    return (
        value.replace("’", "'")
        .replace("‘", "'")
        .replace("`", "'")
        .replace("＇", "'")
        .replace("“", '"')
        .replace("”", '"')
    )


def normalize_language(value: str | None) -> str:
    if not value:
        return "English"
    try:
        region = normalize_region(value)
        return LANGUAGE_BY_REGION[region]
    except SystemExit:
        key = normalize_text(value)
        if key in {"korean", "kr"}:
            return "Korean"
        raise


def normalize_number_part(value: str) -> str:
    value = value.strip().upper()
    if value.isdigit():
        return str(int(value))
    return value


def normalize_card_number(value: str | None) -> str:
    if not value:
        return ""
    value = unicodedata.normalize("NFKD", value)
    value = value.replace("\\", "/").replace("／", "/")
    value = re.sub(r"\s+", "", value.upper())
    if "/" in value:
        return "/".join(normalize_number_part(part) for part in value.split("/", 1))
    return normalize_number_part(value)


def parse_quantity(value: object, row_number: int) -> int:
    if value in (None, ""):
        return 0
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    if not isinstance(value, int):
        raise ValueError(f"Row {row_number}: Quantity must be a whole number, got {value!r}.")
    if value < 0:
        raise ValueError(f"Row {row_number}: Quantity must be 0 or greater, got {value}.")
    return value


def normalize_condition(value: str | None) -> str:
    if not value:
        return "NM/Mint"
    key = normalize_text(value)
    if key in CONDITION_ALIASES:
        return CONDITION_ALIASES[key]
    valid = ", ".join(sorted(set(CONDITION_ALIASES.values())))
    raise ValueError(f"Unknown condition {value!r}. Valid values: {valid}.")


def normalize_holo_pattern(value: str | None) -> str | None:
    key = normalize_text(value)
    if key in HOLO_TYPE_ALIASES:
        return HOLO_TYPE_ALIASES[key] or None
    valid = ", ".join(sorted(value for value in set(HOLO_TYPE_ALIASES.values()) if value))
    raise ValueError(f"Unknown Holo Type {value!r}. Valid values: blank, {valid}.")


def parse_grade_received(value: object, row_number: int) -> float | None:
    if value in (None, ""):
        return None
    if isinstance(value, str):
        value = value.strip()
    try:
        grade = float(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"Row {row_number}: Grade Received must be numeric, got {value!r}.") from exc
    if grade < 1 or grade > 10 or (grade * 2) % 1 != 0:
        raise ValueError(f"Row {row_number}: Grade Received must be 1-10, including half grades.")
    return grade


def header_map(sheet) -> dict[str, int]:
    headers = {}
    for column in range(1, sheet.max_column + 1):
        value = sheet.cell(1, column).value
        if value in (None, ""):
            continue
        header = str(value).strip()
        headers[HEADER_ALIASES.get(header, header)] = column
    missing = set(INPUT_HEADERS) - set(headers)
    if missing:
        raise ValueError(f"Missing required header(s): {', '.join(sorted(missing))}.")
    return headers


def read_intake_rows(path: Path, sheet_name: str) -> list[IntakeRow]:
    workbook = load_workbook(path, data_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Sheet {sheet_name!r} was not found in {path}.")

    sheet = workbook[sheet_name]
    headers = header_map(sheet)
    rows: list[IntakeRow] = []
    for row_number in range(2, sheet.max_row + 1):
        card_name = str(sheet.cell(row_number, headers["Card Name"]).value or "").strip()
        card_number = str(sheet.cell(row_number, headers["Card Number"]).value or "").strip()
        quantity = parse_quantity(sheet.cell(row_number, headers["Quantity"]).value, row_number)
        if not card_name and not card_number and quantity == 0:
            continue
        if not card_name or not card_number:
            raise ValueError(f"Row {row_number}: Card Name and Card Number are required.")

        language = normalize_language(str(sheet.cell(row_number, headers["Language"]).value or "English"))
        condition = normalize_condition(str(sheet.cell(row_number, headers["Condition"]).value or "NM/Mint"))
        holo_type_column = headers.get("Holo Type")
        holo_pattern = (
            normalize_holo_pattern(str(sheet.cell(row_number, holo_type_column).value or ""))
            if holo_type_column
            else None
        )
        grade_company_column = headers.get("Grade Company")
        grade_received_column = headers.get("Grade Received")
        grade_company = (
            str(sheet.cell(row_number, grade_company_column).value or "").strip().upper()
            if grade_company_column
            else ""
        )
        grade_received = (
            parse_grade_received(sheet.cell(row_number, grade_received_column).value, row_number)
            if grade_received_column
            else None
        )
        notes = str(sheet.cell(row_number, headers["Notes"]).value or "").strip() or None
        rows.append(
            IntakeRow(
                row_number=row_number,
                quantity=quantity,
                card_name=card_name,
                card_number=card_number,
                holo_pattern=holo_pattern,
                language=language,
                condition=condition,
                grade_company=grade_company or None,
                grade_received=grade_received,
                notes=notes,
            )
        )
    return rows


def load_existing_cards(session) -> list[ResolvedCard]:
    statement = (
        select(
            Card.id,
            Card.name,
            Card.card_number,
            Card.holo_pattern,
            SetCatalog.set_name,
            SetCatalog.set_code,
            SetCatalog.language,
            Card.tcgcollector_card_id,
            Card.card_detail_url,
        )
        .join(SetCatalog, Card.set_catalog_id == SetCatalog.id, isouter=True)
        .where(Card.game == "Pokemon")
    )
    return [
        ResolvedCard(
            card_id=row.id,
            name=row.name,
            card_number=row.card_number,
            holo_pattern=row.holo_pattern,
            set_name=row.set_name,
            set_code=row.set_code,
            language=row.language,
            tcgcollector_card_id=row.tcgcollector_card_id,
            card_detail_url=row.card_detail_url,
        )
        for row in session.execute(statement)
    ]


def find_existing_match(cards: list[ResolvedCard], intake: IntakeRow) -> ResolvedCard | None:
    target_name = normalize_text(intake.card_name)
    target_number = normalize_card_number(intake.card_number)
    matches = [
        card
        for card in cards
        if normalize_text(card.name) == target_name
        and normalize_card_number(card.card_number) == target_number
        and normalize_text(card.holo_pattern) == normalize_text(intake.holo_pattern)
        and normalize_language(card.language) == intake.language
    ]
    if len(matches) == 1:
        return matches[0]
    return None


def find_default_existing_match(cards: list[ResolvedCard], intake: IntakeRow) -> ResolvedCard | None:
    target_name = normalize_text(intake.card_name)
    target_number = normalize_card_number(intake.card_number)
    matches = [
        card
        for card in cards
        if normalize_text(card.name) == target_name
        and normalize_card_number(card.card_number) == target_number
        and not normalize_text(card.holo_pattern)
        and normalize_language(card.language) == intake.language
    ]
    if len(matches) == 1:
        return matches[0]
    return None


def resolved_card_from_orm(card: Card) -> ResolvedCard:
    set_catalog = card.set_catalog
    return ResolvedCard(
        card_id=card.id,
        name=card.name,
        card_number=card.card_number,
        holo_pattern=card.holo_pattern,
        set_name=set_catalog.set_name if set_catalog else None,
        set_code=set_catalog.set_code if set_catalog else None,
        language=set_catalog.language if set_catalog else None,
        tcgcollector_card_id=card.tcgcollector_card_id,
        card_detail_url=card.card_detail_url,
    )


def create_holo_variant_from_base(session, base_card_id: int, holo_pattern: str) -> ResolvedCard:
    base_card = session.scalar(select(Card).where(Card.id == base_card_id))
    if base_card is None:
        raise ValueError(f"Base card id {base_card_id} was not found.")

    variant = Card(
        set_catalog_id=base_card.set_catalog_id,
        pokedex_id=base_card.pokedex_id,
        name=base_card.name,
        game=base_card.game,
        card_number=base_card.card_number,
        pokemon_name=base_card.pokemon_name,
        rarity=base_card.rarity,
        holo_pattern=holo_pattern,
        source_sequence=base_card.source_sequence,
        tcgcollector_card_id=None,
        card_detail_url=base_card.card_detail_url,
        is_regional_exclusive=base_card.is_regional_exclusive,
        notes=f"synthetic_parallel_variant; holo_pattern={holo_pattern}; base_card_id={base_card.id}",
        primary_image_path=None,
    )
    session.add(variant)
    session.flush()
    return resolved_card_from_orm(variant)


def region_path_for_language(language: str) -> str | None:
    if language == "English":
        return "intl"
    if language == "Japanese":
        return "jp"
    if language == "Simplified Chinese":
        return "cn"
    return None


def region_for_language(language: str) -> str | None:
    if language == "English":
        return "international"
    if language == "Japanese":
        return "japanese"
    if language == "Simplified Chinese":
        return "s-chinese"
    return None


def parse_set_name(full_name: str, card_name: str, card_number: str) -> str | None:
    if "(" not in full_name or ")" not in full_name:
        return None
    inner = full_name.rsplit("(", 1)[1].rsplit(")", 1)[0].strip()
    normalized_inner = normalize_card_number(inner)
    normalized_number = normalize_card_number(card_number)
    if normalized_inner == normalized_number:
        return None

    raw_number = str(card_number).strip()
    for suffix in (raw_number, raw_number.replace(" ", ""), normalized_number):
        if suffix and inner.upper().endswith(suffix.upper()):
            return inner[: -len(suffix)].strip(" -")

    number_pattern = re.escape(raw_number).replace("/", r"\s*/\s*")
    match = re.search(rf"\s+{number_pattern}\s*$", inner, flags=re.IGNORECASE)
    if match:
        return inner[: match.start()].strip(" -")

    if normalize_text(full_name).startswith(normalize_text(card_name)):
        return inner
    return None


def search_tcgcollector_card(intake: IntakeRow) -> tuple[list[RemoteCardMatch], str | None]:
    region_path = region_path_for_language(intake.language)
    if not region_path:
        return [], "unsupported_language"

    url = (
        f"{BASE_URL}/cards/{region_path}"
        f"?cardSearch={quote_plus(normalize_search_query(intake.card_name))}"
        "&releaseDateOrder=newToOld&displayAs=images&cardsPerPage=120"
    )
    html = fetch_url_bytes(url).decode("utf-8", errors="replace")
    target_name = normalize_text(intake.card_name)
    target_number = normalize_card_number(intake.card_number)
    matches: list[RemoteCardMatch] = []
    for card in extract_cards(html):
        if normalize_text(str(card["name"])) != target_name:
            continue
        if normalize_card_number(str(card["card_number"])) != target_number:
            continue
        set_name = parse_set_name(str(card["full_name"]), intake.card_name, intake.card_number)
        if not set_name:
            continue
        matches.append(
            RemoteCardMatch(
                name=str(card["name"]),
                card_number=str(card["card_number"]),
                set_name=set_name,
                tcgcollector_card_id=int(card["tcgcollector_card_id"]),
                card_detail_url=str(card["card_detail_url"]),
            )
        )
    return matches, None


def set_has_cards(conn: sqlite3.Connection, set_catalog_id: int) -> bool:
    row = conn.execute("SELECT COUNT(*) FROM cards WHERE set_catalog_id = ?", (set_catalog_id,)).fetchone()
    return bool(row and row[0])


def import_tcgcollector_set_if_needed(db_path: Path, language: str, set_name: str) -> tuple[bool, str | None]:
    region = region_for_language(language)
    if not region:
        return False, "unsupported_language"

    with sqlite3.connect(db_path) as conn:
        conn.execute("PRAGMA foreign_keys = ON")
        ensure_schema_columns(conn)
        try:
            set_row = resolve_set(conn, region, set_name)
        except SystemExit as exc:
            return False, str(exc)
        if set_has_cards(conn, int(set_row["id"])):
            return False, None

    set_html = fetch_url_bytes(str(set_row["set_url"])).decode("utf-8", errors="replace")
    cards = extract_cards(set_html)
    if not cards:
        return False, f"No cards parsed from {set_row['set_url']}."

    image_dir = image_dir_for_set(set_row, region)
    with sqlite3.connect(db_path) as conn:
        conn.execute("PRAGMA foreign_keys = ON")
        ensure_schema_columns(conn)
        clear_existing_local_images(image_dir)
        with conn:
            for card in cards:
                import_card(conn, set_row, card, image_dir)
    return True, None


def find_set_by_remote_matches(matches: list[RemoteCardMatch]) -> RemoteCardMatch | None:
    unique_by_set = {}
    for match in matches:
        unique_by_set[normalize_text(match.set_name)] = match
    if len(unique_by_set) == 1:
        return next(iter(unique_by_set.values()))
    return None


def pre_import_missing_sets(db_path: Path, intake_rows: list[IntakeRow]) -> None:
    """Import missing TCGcollector sets before inventory writes open a transaction."""

    session_factory = create_session_factory(db_path=db_path)
    with session_factory() as session:
        existing_cards = load_existing_cards(session)
        needed_sets: dict[tuple[str, str], RemoteCardMatch] = {}
        for intake in intake_rows:
            if find_existing_match(existing_cards, intake):
                continue
            if intake.holo_pattern and find_default_existing_match(existing_cards, intake):
                continue
            if intake.language not in SUPPORTED_TCGCOLLECTOR_LANGUAGES:
                continue
            remote_matches, search_status = search_tcgcollector_card(intake)
            if search_status or not remote_matches:
                continue
            remote_match = find_set_by_remote_matches(remote_matches)
            if remote_match:
                needed_sets[(intake.language, remote_match.set_name)] = remote_match

    for language, set_name in sorted(needed_sets):
        _imported, import_error = import_tcgcollector_set_if_needed(db_path, language, set_name)
        if import_error:
            raise RuntimeError(f"Pre-import failed for {language} {set_name}: {import_error}")


def get_or_create_inventory_import(
    session,
    import_label: str,
    import_type: str,
    source_name: str | None = None,
    source_path: str | None = None,
    notes: str | None = None,
) -> InventoryImport:
    import_batch = session.scalar(select(InventoryImport).where(InventoryImport.import_label == import_label))
    if import_batch is None:
        import_batch = InventoryImport(
            import_label=import_label,
            import_type=import_type,
            source_name=source_name,
            source_path=source_path,
            notes=notes,
        )
        session.add(import_batch)
        session.flush()
    return import_batch


def update_inventory_row(
    session,
    intake: IntakeRow,
    card_id: int,
    args: argparse.Namespace,
    import_id: int,
) -> str:
    inventory = session.scalar(select(CardInventory).where(CardInventory.card_id == card_id))
    notes = "; ".join(part for part in [args.notes, intake.notes] if part)
    grading_company_id = None
    if intake.grade_company:
        grading_company_id = session.scalar(
            select(GradingCompany.id).where(GradingCompany.abbreviation == intake.grade_company)
        )
        if grading_company_id is None:
            return f"unknown_grading_company:{intake.grade_company}"

    if not inventory:
        if args.apply:
            session.add(
                CardInventory(
                    card_id=card_id,
                    import_id=import_id,
                    condition=intake.condition,
                    grading_company_id=grading_company_id,
                    grade_received=intake.grade_received,
                    quantity=intake.quantity,
                    cost_basis_cents=0,
                    sale_status="inventory",
                    notes=notes,
                )
            )
        return "inventory_created"

    can_replace = inventory.notes == args.notes or (args.replace_reference and inventory.sale_status == "reference")
    if can_replace:
        if args.apply:
            inventory.import_id = import_id
            inventory.condition = intake.condition
            inventory.grading_company_id = grading_company_id
            inventory.grade_received = intake.grade_received
            inventory.quantity = intake.quantity
            inventory.sale_status = "inventory"
            inventory.notes = notes
        return "inventory_replaced"

    if args.merge_existing:
        if args.apply:
            inventory.import_id = import_id
            inventory.condition = intake.condition or inventory.condition
            inventory.grading_company_id = grading_company_id or inventory.grading_company_id
            inventory.grade_received = intake.grade_received if intake.grade_received is not None else inventory.grade_received
            inventory.quantity = int(inventory.quantity or 0) + intake.quantity
            inventory.sale_status = "inventory"
            inventory.notes = notes if not inventory.notes else f"{inventory.notes}; {notes}"
        return "inventory_merged"

    return "inventory_conflict"


def result_payload(
    card: ResolvedCard | None,
    status: str,
    remote_match: RemoteCardMatch | None = None,
) -> dict[str, object | None]:
    return {
        "Matched Card ID": card.card_id if card else None,
        "Resolution Status": status,
        "Matched Set": card.set_name if card else remote_match.set_name if remote_match else None,
        "Matched Set Code": card.set_code if card else None,
        "Matched Holo Type": card.holo_pattern if card else None,
        "TCGcollector Card ID": card.tcgcollector_card_id if card else remote_match.tcgcollector_card_id if remote_match else None,
        "Card Detail URL": card.card_detail_url if card else remote_match.card_detail_url if remote_match else None,
    }


def write_results_to_workbook(path: Path, sheet_name: str, results: dict[int, dict[str, object | None]]) -> None:
    workbook = load_workbook(path)
    sheet = workbook[sheet_name]
    headers = header_map(sheet)
    next_column = sheet.max_column + 1
    for header in RESULT_HEADERS:
        if header not in headers:
            headers[header] = next_column
            sheet.cell(1, next_column, header)
            next_column += 1
    for row_number, row_result in results.items():
        for header, value in row_result.items():
            sheet.cell(row_number, headers[header], value)
    workbook.save(path)


def main() -> None:
    args = parse_args()
    db_path = resolve_path(args.db_path)
    workbook_path = resolve_path(args.workbook)
    if not db_path.exists():
        raise SystemExit(f"Database not found: {db_path}. Run scripts/init_db.py first.")
    if not workbook_path.exists():
        raise SystemExit(f"Workbook not found: {workbook_path}")

    intake_rows = read_intake_rows(workbook_path, args.sheet)
    if args.apply and args.import_missing_sets:
        pre_import_missing_sets(db_path, intake_rows)

    session_factory = create_session_factory(db_path=db_path)
    results: dict[int, dict[str, object | None]] = {}
    summary: dict[str, int] = {}

    with session_factory() as session:
        import_label = args.import_label.strip() or workbook_path.stem
        import_batch = get_or_create_inventory_import(
            session,
            import_label=import_label,
            import_type="binder",
            source_name="binder_intake_workbook",
            source_path=str(workbook_path),
            notes=args.notes,
        )
        existing_cards = load_existing_cards(session)
        for intake in intake_rows:
            status = ""
            remote_match = None
            card = find_existing_match(existing_cards, intake)
            if card:
                status = "matched_existing"
            elif intake.holo_pattern:
                base_card = find_default_existing_match(existing_cards, intake)
                if base_card and args.apply:
                    card = create_holo_variant_from_base(session, base_card.card_id, intake.holo_pattern)
                    existing_cards.append(card)
                    status = f"created_holo_variant:{intake.holo_pattern}"
                elif base_card:
                    status = f"holo_variant_ready_to_create:{intake.holo_pattern}"
            elif intake.language not in SUPPORTED_TCGCOLLECTOR_LANGUAGES:
                status = "unsupported_language"

            if not card and not status:
                remote_matches, search_status = search_tcgcollector_card(intake)
                remote_match = find_set_by_remote_matches(remote_matches)
                if search_status:
                    status = search_status
                elif not remote_matches:
                    status = "unresolved_no_remote_match"
                elif not remote_match:
                    status = "ambiguous_remote_match"
                elif not args.import_missing_sets:
                    status = f"remote_match_set_not_imported: {remote_match.set_name}"
                elif not args.apply:
                    status = f"remote_match_set_ready_to_import: {remote_match.set_name}"
                else:
                    imported, import_error = import_tcgcollector_set_if_needed(db_path, intake.language, remote_match.set_name)
                    if import_error:
                        status = f"set_import_failed: {import_error}"
                    else:
                        session.expire_all()
                        existing_cards = load_existing_cards(session)
                        card = find_existing_match(existing_cards, intake)
                        status = "matched_after_set_import" if imported else "matched_after_existing_set"
                        if not card and intake.holo_pattern:
                            base_card = find_default_existing_match(existing_cards, intake)
                            if base_card and args.apply:
                                card = create_holo_variant_from_base(session, base_card.card_id, intake.holo_pattern)
                                existing_cards.append(card)
                                status = f"{status}; created_holo_variant:{intake.holo_pattern}"
                            elif base_card:
                                status = f"{status}; holo_variant_ready_to_create:{intake.holo_pattern}"
                        if not card and "holo_variant_ready_to_create" not in status:
                            status = "set_imported_but_card_not_matched"

            if not card and intake.holo_pattern and not status:
                status = f"holo_variant_unresolved:{intake.holo_pattern}"

            if card and intake.quantity > 0:
                inventory_status = update_inventory_row(session, intake, card.card_id, args, import_batch.id)
                status = f"{status}; {inventory_status}"

            results[intake.row_number] = result_payload(card, status, remote_match)
            summary[status.split(";", 1)[0]] = summary.get(status.split(";", 1)[0], 0) + 1

        if args.apply:
            session.commit()
        else:
            session.rollback()

    if args.write_results:
        write_results_to_workbook(workbook_path, args.sheet, results)

    action = "Applied" if args.apply else "Dry run"
    print(f"{action} binder import for {len(intake_rows)} workbook rows.")
    for status, count in sorted(summary.items()):
        print(f"{status}: {count}")
    if args.write_results:
        print(f"Wrote resolution results to {workbook_path}")


if __name__ == "__main__":
    main()
