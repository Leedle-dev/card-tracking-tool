from pathlib import Path
import argparse
import re
import sqlite3

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / "data" / "card_tracker.sqlite"
DEFAULT_OUTPUT_DIR = ROOT / "reports" / "inventory"


def slugify(value: str) -> str:
    slug = re.sub(r"[^a-z0-9]+", "_", value.lower()).strip("_")
    return slug or "inventory"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build a protected inventory-count workbook for an imported card set."
    )
    # Argument examples:
    #   --set-code CBB5C
    #   --set-code CBB5C --output reports/inventory/gem_pack_vol_5_inventory.xlsx
    #   --set-name "Gem Pack Vol. 5" --language s-chinese
    parser.add_argument("--set-code", default="", help="Set code to export, such as CBB5C.")
    parser.add_argument("--set-name", default="", help="Set name to export if no set code is provided.")
    parser.add_argument("--language", default="", help="Optional language filter, such as s-chinese.")
    parser.add_argument("--output", default="", help="Optional XLSX output path.")
    return parser.parse_args()


def fetch_cards(args: argparse.Namespace) -> list[sqlite3.Row]:
    if not args.set_code and not args.set_name:
        raise SystemExit("Provide --set-code or --set-name.")

    where = []
    params: list[str] = []

    if args.set_code:
        where.append("LOWER(sc.set_code) = LOWER(?)")
        params.append(args.set_code)

    if args.set_name:
        where.append("LOWER(sc.set_name) = LOWER(?)")
        params.append(args.set_name)

    if args.language:
        where.append("LOWER(sc.language) = LOWER(?)")
        params.append(args.language)

    query = f"""
        SELECT
            c.id AS card_id,
            c.set_catalog_id,
            sc.tcgcollector_set_id,
            c.tcgcollector_card_id,
            sc.language,
            sc.release_year,
            sc.release_date_text,
            sc.set_code,
            sc.set_name,
            c.card_number,
            c.name AS card_name,
            c.pokemon_name,
            c.rarity,
            c.holo_pattern,
            c.source_sequence,
            c.is_regional_exclusive,
            c.primary_image_path,
            COALESCE(SUM(CASE WHEN ci.sale_status = 'inventory' THEN ci.quantity ELSE 0 END), 0)
                AS current_db_inventory_count
        FROM cards c
        JOIN set_catalog sc ON sc.id = c.set_catalog_id
        LEFT JOIN card_inventory ci ON ci.card_id = c.id
        WHERE {" AND ".join(where)}
        GROUP BY c.id
        ORDER BY sc.language, sc.set_name, c.source_sequence, c.card_number, c.id
    """

    with sqlite3.connect(DB_PATH) as conn:
        conn.row_factory = sqlite3.Row
        return conn.execute(query, params).fetchall()


def build_workbook(rows: list[sqlite3.Row], output_path: Path) -> None:
    if not rows:
        raise SystemExit("No cards found for the requested set.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventory"

    set_name = rows[0]["set_name"]
    set_code = rows[0]["set_code"]
    language = rows[0]["language"]

    columns = [
        ("inventory_count", "Count"),
        ("card_id", "Card ID"),
        ("set_catalog_id", "Set Catalog ID"),
        ("tcgcollector_set_id", "TCGcollector Set ID"),
        ("tcgcollector_card_id", "TCGcollector Card ID"),
        ("language", "Language"),
        ("release_year", "Release Year"),
        ("release_date_text", "Release Date"),
        ("set_code", "Set Code"),
        ("set_name", "Set Name"),
        ("card_number", "Card Number"),
        ("card_name", "Card Name"),
        ("pokemon_name", "Pokemon Name"),
        ("rarity", "Rarity"),
        ("holo_pattern", "Holo Pattern"),
        ("source_sequence", "Source Sequence"),
        ("is_regional_exclusive", "Regional Exclusive"),
        ("primary_image_path", "Image Path"),
    ]

    ws["A1"] = f"{set_name} ({set_code}) Inventory"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F4E78")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))

    ws["A2"] = "Only the Count column is unlocked for editing. All other columns are stable keys for future database ingestion."
    ws["A2"].font = Font(italic=True, color="555555")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(columns))

    header_row = 4
    for col_idx, (_, header) in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, row in enumerate(rows, start=header_row + 1):
        values = {
            "inventory_count": row["current_db_inventory_count"],
            "is_regional_exclusive": "Yes" if row["is_regional_exclusive"] else "No",
        }
        values.update(dict(row))

        for col_idx, (key, _) in enumerate(columns, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=values.get(key))
            cell.alignment = Alignment(vertical="top", wrap_text=key in {"card_name", "primary_image_path"})

        count_cell = ws.cell(row=row_idx, column=1)
        count_cell.protection = Protection(locked=False)
        count_cell.fill = PatternFill("solid", fgColor="FFF2CC")
        count_cell.alignment = Alignment(horizontal="center", vertical="center")

    last_row = header_row + len(rows)
    table_ref = f"A{header_row}:R{last_row}"
    table = Table(displayName="GemPackInventory", ref=table_ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)

    validation = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=False)
    validation.error = "Inventory count must be a whole number of 0 or greater."
    validation.errorTitle = "Invalid inventory count"
    validation.prompt = "Enter owned quantity for this card."
    validation.promptTitle = "Inventory count"
    ws.add_data_validation(validation)
    validation.add(f"A{header_row + 1}:A{last_row}")

    ws["A4"].comment = Comment(
        "Editable inventory quantity. Future ingestion can key each count by Card ID.",
        "Codex",
    )

    ws.freeze_panes = "A5"
    ws.auto_filter.ref = table_ref

    widths = {
        "A": 12,
        "B": 10,
        "C": 14,
        "D": 18,
        "E": 20,
        "F": 14,
        "G": 12,
        "H": 16,
        "I": 12,
        "J": 22,
        "K": 14,
        "L": 28,
        "M": 18,
        "N": 20,
        "O": 22,
        "P": 15,
        "Q": 18,
        "R": 60,
    }
    for col_letter, width in widths.items():
        ws.column_dimensions[col_letter].width = width

    for row_number in range(header_row + 1, last_row + 1):
        ws.row_dimensions[row_number].height = 24

    ws.protection.sheet = True
    ws.protection.objects = True
    ws.protection.scenarios = True
    ws.protection.selectLockedCells = False
    ws.protection.selectUnlockedCells = False

    meta = wb.create_sheet("Ingestion_Metadata")
    meta_rows = [
        ("template_name", "card_inventory_count_template"),
        ("template_version", "1"),
        ("set_code", set_code),
        ("set_name", set_name),
        ("language", language),
        ("record_count", len(rows)),
        ("editable_columns", "inventory_count"),
        ("ingestion_key", "card_id"),
    ]
    for idx, (key, value) in enumerate(meta_rows, start=1):
        meta.cell(row=idx, column=1, value=key).font = Font(bold=True)
        meta.cell(row=idx, column=2, value=value)
    meta.column_dimensions["A"].width = 24
    meta.column_dimensions["B"].width = 44
    meta.protection.sheet = True

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def main() -> None:
    args = parse_args()
    rows = fetch_cards(args)

    if args.output:
        output_path = Path(args.output)
        if not output_path.is_absolute():
            output_path = ROOT / output_path
    else:
        first = rows[0]
        output_path = DEFAULT_OUTPUT_DIR / (
            f"{slugify(first['set_code'])}_{slugify(first['set_name'])}_inventory.xlsx"
        )

    build_workbook(rows, output_path)
    print(output_path)


if __name__ == "__main__":
    main()
