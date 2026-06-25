from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence
import csv
import json
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from card_tracker.report_paths import PROJECT_ROOT, slugify_report_label, timestamped_report_dir


TableRow = Mapping[str, Any]


@dataclass
class SheetSpec:
    name: str
    rows: Sequence[TableRow] = field(default_factory=list)
    fieldnames: Sequence[str] | None = None
    title: str | None = None
    freeze_panes: str | None = "A2"
    auto_filter: bool = True
    table_name: str | None = None
    table_style: str = "TableStyleMedium2"
    header_fill: str = "1F2937"
    header_font_color: str = "FFFFFF"
    header_bold: bool = True
    wrap_text: bool = True
    column_widths: Mapping[str | int, float] = field(default_factory=dict)
    column_fills: Mapping[str | int, str] = field(default_factory=dict)
    number_formats: Mapping[str | int, str] = field(default_factory=dict)
    hidden_columns: Sequence[str | int] = field(default_factory=list)
    currency_columns: Sequence[str | int] = field(default_factory=list)
    percent_columns: Sequence[str | int] = field(default_factory=list)
    integer_columns: Sequence[str | int] = field(default_factory=list)
    max_auto_width: int = 42


@dataclass
class WorkbookSpec:
    sheets: Sequence[SheetSpec]
    filename: str
    creator: str = "Card Tracking Tool"
    subject: str = "Generated report"


class ReportBuilder:
    """Small helper for timestamped report folders and common report files."""

    def __init__(
        self,
        report_type: str,
        label: str | None = None,
        output_dir: Path | str | None = None,
        metadata: Mapping[str, Any] | None = None,
    ) -> None:
        self.report_type = slugify_report_label(report_type)
        self.label = slugify_report_label(label, self.report_type)
        self.created_at = datetime.now().astimezone().isoformat(timespec="seconds")
        if output_dir is None:
            self.output_dir = timestamped_report_dir(self.report_type, self.label)
        else:
            output_path = Path(output_dir)
            self.output_dir = output_path if output_path.is_absolute() else PROJECT_ROOT / output_path
            self.output_dir.mkdir(parents=True, exist_ok=True)
        self.raw_dir = self.output_dir / "raw"
        self.files: dict[str, str] = {}
        self.metadata: dict[str, Any] = {
            "report_type": self.report_type,
            "label": self.label,
            "created_at": self.created_at,
        }
        if metadata:
            self.metadata.update(dict(metadata))

    def path(self, filename: str, subdir: str | None = None) -> Path:
        base = self.output_dir / subdir if subdir else self.output_dir
        base.mkdir(parents=True, exist_ok=True)
        return base / filename

    def raw_path(self, filename: str) -> Path:
        self.raw_dir.mkdir(parents=True, exist_ok=True)
        return self.raw_dir / filename

    def record_file(self, key: str, path: Path) -> Path:
        self.files[key] = str(path)
        return path

    def write_text(self, filename: str, text: str, key: str | None = None, subdir: str | None = None) -> Path:
        path = self.path(filename, subdir=subdir)
        path.write_text(text, encoding="utf-8")
        return self.record_file(key or filename, path)

    def write_json(
        self,
        filename: str,
        data: Any,
        key: str | None = None,
        raw: bool = False,
        subdir: str | None = None,
    ) -> Path:
        path = self.raw_path(filename) if raw else self.path(filename, subdir=subdir)
        path.write_text(json.dumps(data, indent=2, sort_keys=True, default=str), encoding="utf-8")
        return self.record_file(key or filename, path)

    def write_tsv(
        self,
        filename: str,
        rows: Sequence[TableRow],
        fieldnames: Sequence[str] | None = None,
        key: str | None = None,
    ) -> Path:
        return self.write_delimited(filename, rows, delimiter="\t", fieldnames=fieldnames, key=key)

    def write_csv(
        self,
        filename: str,
        rows: Sequence[TableRow],
        fieldnames: Sequence[str] | None = None,
        key: str | None = None,
    ) -> Path:
        return self.write_delimited(filename, rows, delimiter=",", fieldnames=fieldnames, key=key)

    def write_delimited(
        self,
        filename: str,
        rows: Sequence[TableRow],
        delimiter: str,
        fieldnames: Sequence[str] | None = None,
        key: str | None = None,
    ) -> Path:
        path = self.path(filename)
        headers = list(fieldnames or infer_fieldnames(rows))
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=headers, delimiter=delimiter, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(rows)
        return self.record_file(key or filename, path)

    def write_workbook(self, spec: WorkbookSpec, key: str | None = None) -> Path:
        path = self.path(spec.filename)
        workbook = Workbook()
        workbook.properties.creator = spec.creator
        workbook.properties.subject = spec.subject

        default_sheet = workbook.active
        workbook.remove(default_sheet)
        for sheet in spec.sheets:
            write_sheet(workbook, sheet)

        workbook.save(path)
        return self.record_file(key or spec.filename, path)

    def write_metadata(self, filename: str = "report_metadata.json") -> Path:
        metadata = {**self.metadata, "output_dir": str(self.output_dir), "files": self.files}
        return self.write_json(filename, metadata, key="metadata")


def infer_fieldnames(rows: Sequence[TableRow]) -> list[str]:
    if not rows:
        return []
    seen: list[str] = []
    for row in rows:
        for key in row.keys():
            if key not in seen:
                seen.append(str(key))
    return seen


def resolve_column_key(column: str | int, fieldnames: Sequence[str]) -> int:
    if isinstance(column, int):
        return column
    if column in fieldnames:
        return list(fieldnames).index(column) + 1
    if re.fullmatch(r"[A-Za-z]+", column):
        result = 0
        for char in column.upper():
            result = result * 26 + ord(char) - ord("A") + 1
        return result
    raise KeyError(f"Unknown column reference: {column}")


def write_sheet(workbook: Workbook, spec: SheetSpec) -> None:
    fieldnames = list(spec.fieldnames or infer_fieldnames(spec.rows))
    sheet = workbook.create_sheet(title=safe_sheet_name(spec.name))

    header_row = 1
    if spec.title:
        sheet.cell(1, 1, spec.title)
        sheet.cell(1, 1).font = Font(bold=True, size=14)
        header_row = 3

    for column_index, header in enumerate(fieldnames, start=1):
        cell = sheet.cell(header_row, column_index, header)
        cell.font = Font(bold=spec.header_bold, color=spec.header_font_color)
        cell.fill = PatternFill("solid", fgColor=spec.header_fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, row in enumerate(spec.rows, start=header_row + 1):
        for column_index, header in enumerate(fieldnames, start=1):
            cell = sheet.cell(row_index, column_index, row.get(header))
            if spec.wrap_text:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    if spec.freeze_panes:
        sheet.freeze_panes = spec.freeze_panes

    last_row = max(header_row + len(spec.rows), header_row)
    last_column = max(len(fieldnames), 1)
    table_range = f"A{header_row}:{get_column_letter(last_column)}{last_row}"
    if spec.auto_filter:
        sheet.auto_filter.ref = table_range
    if spec.table_name and spec.rows and fieldnames:
        table = Table(displayName=safe_table_name(spec.table_name), ref=table_range)
        table.tableStyleInfo = TableStyleInfo(name=spec.table_style, showRowStripes=True, showColumnStripes=False)
        sheet.add_table(table)

    apply_column_formatting(sheet, spec, fieldnames, header_row, last_row)


def apply_column_formatting(sheet: Any, spec: SheetSpec, fieldnames: Sequence[str], header_row: int, last_row: int) -> None:
    explicit_formats: dict[int, str] = {}
    for column in spec.currency_columns:
        explicit_formats[resolve_column_key(column, fieldnames)] = '$#,##0.00'
    for column in spec.percent_columns:
        explicit_formats[resolve_column_key(column, fieldnames)] = '0.00%'
    for column in spec.integer_columns:
        explicit_formats[resolve_column_key(column, fieldnames)] = '0'
    for column, number_format in spec.number_formats.items():
        explicit_formats[resolve_column_key(column, fieldnames)] = number_format

    for column_index, number_format in explicit_formats.items():
        for row_index in range(header_row + 1, last_row + 1):
            sheet.cell(row_index, column_index).number_format = number_format

    for column, fill_color in spec.column_fills.items():
        column_index = resolve_column_key(column, fieldnames)
        fill = PatternFill("solid", fgColor=fill_color)
        for row_index in range(header_row + 1, last_row + 1):
            sheet.cell(row_index, column_index).fill = fill

    for column in spec.hidden_columns:
        column_letter = get_column_letter(resolve_column_key(column, fieldnames))
        sheet.column_dimensions[column_letter].hidden = True

    widths = {resolve_column_key(column, fieldnames): width for column, width in spec.column_widths.items()}
    for column_index, header in enumerate(fieldnames, start=1):
        if column_index in widths:
            width = widths[column_index]
        else:
            sample_values = [str(sheet.cell(row, column_index).value or "") for row in range(header_row, last_row + 1)]
            width = min(spec.max_auto_width, max(10, max(len(value) for value in sample_values) + 2))
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def safe_sheet_name(name: str) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", name).strip() or "Sheet"
    return cleaned[:31]


def safe_table_name(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9_]", "_", name).strip("_")
    if not cleaned or cleaned[0].isdigit():
        cleaned = f"Table_{cleaned or 'Report'}"
    return cleaned[:255]


def read_table_file(path: Path | str, delimiter: str | None = None) -> list[dict[str, str]]:
    table_path = Path(path)
    if not table_path.is_absolute():
        table_path = PROJECT_ROOT / table_path
    if delimiter is None:
        delimiter = "\t" if table_path.suffix.lower() == ".tsv" else ","
    with table_path.open("r", newline="", encoding="utf-8-sig") as handle:
        return list(csv.DictReader(handle, delimiter=delimiter))


def coerce_rows(rows: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
    return [dict(row) for row in rows]
